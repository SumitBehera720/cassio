import openpyxl
import sqlite3
import sys

# Ensure UTF-8
sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('Watch_Product_Specifications (3).xlsx')
sheet = wb.active
headers = [sheet.cell(1, c).value for c in range(1, sheet.max_column + 1)]

# Map Excel row to Product ID
row_to_pid = {
    2: 25, # Tommy Hilfiger
    3: 26, # Tissot 1853 Chronograph
    4: 27, # Fossil Automatic Skeleton (Leather)
    5: 28, # Fossil Automatic Skeleton (Metal)
    6: 29, # Rado Automatic Skeleton
}

product_meta = {
    25: {
        'name': 'Tommy Hilfiger Lorenzo TH1792117',
        'slug': 'tommy-hilfiger-lorenzo-th1792117',
        'subtitle': 'Analog Multifunction · Deep Green Sunray Dial · Silver Stainless Steel',
        'sku': '1792117000',
        'price': 2500,
        'compare_at': 15999,
        'intro': 'The Tommy Hilfiger Lorenzo (TH1792117) is an iconic multifunction timepiece combining nautical-inspired prestige with everyday elegance. Showcasing an eye-catching deep sunray green dial with a high-contrast black rotating bezel and a solid stainless steel bracelet.'
    },
    26: {
        'name': 'Tissot 1853 Chronograph',
        'slug': 'tissot-1853-chronograph',
        'subtitle': 'Analog Chronograph · 3-Subdial Glossy Black Dial · Crocodile-Pattern Leather',
        'sku': 'T1853-CHRONO-BLK',
        'price': 1800,
        'compare_at': 4999,
        'intro': 'Swiss heritage meets motorsport precision. The Tissot 1853 Chronograph features a deep glossy black dial with 3 functioning subdials, a silver-tone tachymeter outer bezel, date display, and a luxurious crocodile-embossed genuine leather strap.'
    },
    27: {
        'name': 'Fossil Automatic Skeleton Watch (Leather Strap)',
        'slug': 'fossil-automatic-skeleton-leather',
        'subtitle': 'Self-Winding Skeleton · Smoke Dial · Dark Brown Textured Leather Strap',
        'sku': 'FSL-AUTO-SKEL-BRN',
        'price': 3000,
        'compare_at': 9999,
        'intro': 'Pure mechanical fascination. This Fossil Automatic timepiece reveals the intricate mechanical heartbeat of self-winding horology through an open-heart smoke-tinted skeleton dial, accented by a knurled black bezel and rich coffee-brown leather.'
    },
    28: {
        'name': 'Fossil Automatic Skeleton Watch (Metal Bracelet)',
        'slug': 'fossil-automatic-skeleton-metal',
        'subtitle': 'Self-Winding Skeleton · Roman XII Dial · Gunmetal Solid Steel Bracelet',
        'sku': 'FSL-AUTO-SKEL-MET',
        'price': 3000,
        'compare_at': 12000,
        'intro': 'Bold, industrial luxury. Crafted from solid gunmetal black stainless steel, this Fossil Automatic Skeleton watch features prominent Roman numerals, an exposed gear-train escapement, sub-seconds display, and a heavy-gauge multi-link bracelet.'
    },
    29: {
        'name': 'Rado Automatic Skeleton Watch',
        'slug': 'rado-automatic-skeleton',
        'subtitle': 'Self-Winding Skeleton · Blue & Black Dial · Premium Solid Steel Bracelet',
        'sku': 'RADO-AUTO-SKEL-BLU',
        'price': 4000,
        'compare_at': 16000,
        'intro': 'An architectural masterpiece on the wrist. The Rado Automatic Skeleton features a captivating two-tone blue and black skeleton dial with golden mechanical gears, housed in an ultra-sleek black stainless steel case with integrated link bracelet.'
    }
}

conn = sqlite3.connect('classio/database/database.sqlite')
c = conn.cursor()

for r, pid in row_to_pid.items():
    row_data = {}
    for col_idx, h in enumerate(headers):
        val = sheet.cell(r, col_idx + 1).value
        if val is not None and str(val).strip() != '':
            row_data[h] = str(val).strip()

    meta = product_meta[pid]
    
    # Pricing info
    orig_price = row_data.get('Original / Listed Price', f"₹{meta['compare_at']:,}")
    disc_50 = row_data.get('50% Discount Price', '')
    spec_disc = row_data.get('Special Extra Discount', '')
    final_price_str = row_data.get('Final Sale Price', f"₹{meta['price']:,}")
    tot_disc = row_data.get('Total Discount', '')
    tot_pct = row_data.get('Total Discount Percentage', '')
    you_save = row_data.get('You Save', '')
    final_offer = row_data.get('Final Offer', f"₹{meta['price']:,}/- ONLY")

    # Build Price Breakdown HTML Card
    price_html = f'''
<div class="mb-5 bg-gradient-to-br from-amber-50 to-orange-50/40 border border-amber-200/80 rounded-xl p-4 text-xs font-sans">
    <div class="flex items-center justify-between border-b border-amber-200/60 pb-2 mb-2.5">
        <span class="font-bold text-amber-900 uppercase tracking-wider text-[11px] flex items-center gap-1.5">
            <svg class="w-3.5 h-3.5 text-amber-700" fill="currentColor" viewBox="0 0 20 20"><path fill-rule="evenodd" d="M10 18a8 8 0 100-16 8 8 0 000 16zm1-12a1 1 0 10-2 0v4a1 1 0 00.293.707l2.828 2.829a1 1 0 101.415-1.415L11 9.586V6z" clip-rule="evenodd"></path></svg>
            Special Offer Breakdown
        </span>
        <span class="bg-[#B91C1C] text-white text-[10px] font-bold px-2 py-0.5 rounded-full">{tot_pct if tot_pct else 'SPECIAL SALE'}</span>
    </div>
    <div class="grid grid-cols-2 sm:grid-cols-3 gap-2.5 text-gray-700">
        <div>
            <div class="text-[10px] text-gray-500 uppercase tracking-wide">Listed Price (MRP)</div>
            <div class="font-semibold text-gray-500 line-through">{orig_price}</div>
        </div>
        {f'<div><div class="text-[10px] text-gray-500 uppercase tracking-wide">Special Discount</div><div class="font-semibold text-gray-700">{spec_disc}</div></div>' if spec_disc else ''}
        <div>
            <div class="text-[10px] text-emerald-700 font-bold uppercase tracking-wide">You Save</div>
            <div class="font-bold text-emerald-600">{you_save if you_save else (tot_disc if tot_disc else 'Huge Savings')}</div>
        </div>
        <div class="col-span-2 sm:col-span-3 pt-2 border-t border-amber-200/50 flex items-center justify-between">
            <span class="text-xs font-bold text-gray-900">Final Classio Offer Price:</span>
            <span class="text-sm font-extrabold text-[#0B1A2C] bg-white px-2.5 py-1 rounded-md border border-amber-300 shadow-xs">{final_offer if final_offer else final_price_str}</span>
        </div>
    </div>
</div>
'''

    # Build Technical Specifications Table
    specs_keys = [
        ('Brand', 'Brand'),
        ('Collection', 'Collection'),
        ('Model', 'Model Number'),
        ('Reference', 'Reference'),
        ('SKU', 'SKU'),
        ('Gender', 'Gender'),
        ('Watch Type', 'Watch Type'),
        ('Movement', 'Movement'),
        ('Dial', 'Dial'),
        ('Dial Type', 'Dial Type'),
        ('Bezel', 'Bezel'),
        ('Case', 'Case Material'),
        ('Case Shape', 'Case Shape'),
        ('Case Size', 'Case Size'),
        ('Case Thickness', 'Case Thickness'),
        ('Crystal', 'Crystal Glass'),
        ('Bracelet/Strap', 'Bracelet / Strap'),
        ('Color', 'Color'),
        ('Design', 'Design'),
        ('Clasp', 'Clasp Type'),
        ('Functions', 'Functions'),
        ('Power Reserve', 'Power Reserve'),
        ('Water Resistance', 'Water Resistance'),
        ('Style', 'Style'),
    ]

    table_rows = []
    row_count = 0
    for excel_k, label in specs_keys:
        if excel_k in row_data and row_data[excel_k]:
            bg_cls = 'bg-white' if row_count % 2 == 0 else 'bg-gray-50/60'
            val = row_data[excel_k]
            table_rows.append(f"<tr class='{bg_cls}'><td class='px-3.5 py-2 text-gray-500 font-medium w-2/5 sm:w-1/3 border-r border-gray-100'>{label}</td><td class='px-3.5 py-2 text-gray-900 font-semibold'>{val}</td></tr>")
            row_count += 1

    table_html = f'''
<div class="border border-gray-200 rounded-xl overflow-hidden mt-4">
    <div class="bg-gray-100/80 px-4 py-2.5 border-b border-gray-200 font-sans text-xs font-bold uppercase tracking-wider text-gray-800 flex items-center justify-between">
        <span>Technical Specifications</span>
        <span class="text-[10px] text-gray-500 font-normal normal-case">Verified Manufacturer Specs</span>
    </div>
    <table class="w-full text-xs font-sans divide-y divide-gray-100">
        <tbody class="divide-y divide-gray-100">
            {''.join(table_rows)}
        </tbody>
    </table>
</div>
'''

    full_description = f"<p class='text-sm text-gray-700 leading-relaxed mb-4'>{meta['intro']}</p>{price_html}{table_html}"

    # Update products table
    c.execute("""
        UPDATE products 
        SET name = ?, slug = ?, subtitle = ?, description = ?
        WHERE id = ?
    """, (meta['name'], meta['slug'], meta['subtitle'], full_description, pid))

    # Update product_variants table (set size=null, color=null so apparel size pickers don't appear)
    c.execute("""
        UPDATE product_variants
        SET price = ?, compare_at_price = ?, sku = ?, size = NULL, color = NULL
        WHERE product_id = ?
    """, (meta['price'], meta['compare_at'], meta['sku'], pid))

    print(f"Updated PID {pid} ({meta['name']}): Price=₹{meta['price']}, CompareAt=₹{meta['compare_at']}, SKU={meta['sku']}")

conn.commit()
conn.close()
print("All watch specifications successfully synchronized!")
