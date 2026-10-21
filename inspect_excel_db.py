import openpyxl
import json
import sqlite3
import os
import sys

# Set stdout encoding
sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('Watch_Product_Specifications (3).xlsx')
sheet = wb.active
headers = [sheet.cell(1, c).value for c in range(1, sheet.max_column + 1)]

print("HEADERS:")
for i, h in enumerate(headers):
    print(f"  Col {i+1}: {h}")

rows_data = []
for r in range(2, sheet.max_row + 1):
    row_dict = {}
    for c in range(1, sheet.max_column + 1):
        h = headers[c-1]
        val = sheet.cell(r, c).value
        row_dict[h] = str(val) if val is not None else ""
    rows_data.append(row_dict)

print("\n--- EXCEL ROWS ---")
for idx, row in enumerate(rows_data):
    print(f"\n--- Row {idx+2}: {row.get('Brand')} - {row.get('Model')} ---")
    for k, v in row.items():
        if v:
            print(f"  {k}: {v}")

print("\n--- PRODUCTS IN CATEGORY 3 OR WITH WATCH IN NAME ---")
conn = sqlite3.connect('classio/database/database.sqlite')
cursor = conn.cursor()
cursor.execute("SELECT id, category_id, subcategory_id, name, slug FROM products WHERE category_id = 3 OR lower(name) LIKE '%watch%'")
for r in cursor.fetchall():
    print(r)
conn.close()
